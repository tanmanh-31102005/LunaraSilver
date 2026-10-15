import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

OUTPUT_DOCX = os.path.abspath("docs/final/Bao_cao_Do_An_Lunara_Silver_Hoan_Chinh.docx")
ASSET_DIR = os.path.abspath("docs/final/assets")

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.bold = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(26, 28, 44) # Midnight color
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(33, 37, 41)
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(13)
        run.italic = True
        run.font.color.rgb = RGBColor(73, 80, 87)
    return p

def add_body(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13)
        r_pre.bold = True
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    if italic:
        r.italic = True
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13)
        r_pre.bold = True
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    return p

def add_image_with_caption(doc, img_rel_path, caption_text, width_inches=6.2, max_height_inches=7.4):
    img_full_path = os.path.join(ASSET_DIR, img_rel_path)
    if not os.path.exists(img_full_path):
        print(f"Warning: Image {img_full_path} not found!")
        return
        
    final_width = width_inches
    try:
        from PIL import Image
        with Image.open(img_full_path) as im:
            w_px, h_px = im.size
            if h_px > 0:
                aspect = w_px / h_px
                # If image is very tall, scale width down so height <= max_height_inches
                calc_height = width_inches / aspect
                if calc_height > max_height_inches:
                    final_width = max_height_inches * aspect
    except Exception as e:
        final_width = width_inches

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True  # Always keep picture with caption!
    run_img = p_img.add_run()
    run_img.add_picture(img_full_path, width=Inches(final_width))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.paragraph_format.keep_with_next = False
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = 'Times New Roman'
    run_cap.font.size = Pt(11)
    run_cap.italic = True
    run_cap.font.color.rgb = RGBColor(100, 100, 100)

def add_landscape_diagram_page(doc, img_rel_path, caption_text, width_inches=10.2, max_height_inches=6.6):
    """
    Creates a dedicated A4 Landscape page for a complex diagram:
    Section Break (Landscape) -> Image + Caption -> Section Break (Portrait).
    """
    from docx.enum.section import WD_ORIENT, WD_SECTION
    
    # 1. New Landscape Section
    sec_land = doc.add_section(WD_SECTION.NEW_PAGE)
    sec_land.orientation = WD_ORIENT.LANDSCAPE
    sec_land.page_width = Inches(11.69)
    sec_land.page_height = Inches(8.27)
    sec_land.top_margin = Inches(0.55)
    sec_land.bottom_margin = Inches(0.55)
    sec_land.left_margin = Inches(0.60)
    sec_land.right_margin = Inches(0.60)
    
    # Insert Image + Caption (max width 10.2 in, max height 6.6 in)
    add_image_with_caption(doc, img_rel_path, caption_text, width_inches=width_inches, max_height_inches=max_height_inches)
    
    # 2. Return to Standard Academic Portrait Section
    sec_port = doc.add_section(WD_SECTION.NEW_PAGE)
    sec_port.orientation = WD_ORIENT.PORTRAIT
    sec_port.page_width = Inches(8.27)
    sec_port.page_height = Inches(11.69)
    sec_port.left_margin = Inches(1.18)    # 3.0 cm
    sec_port.right_margin = Inches(0.79)   # 2.0 cm
    sec_port.top_margin = Inches(0.98)     # 2.5 cm
    sec_port.bottom_margin = Inches(0.98)  # 2.5 cm


def add_code_block(doc, code_text, caption=""):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(3)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.bold = True
        run_cap.font.color.rgb = RGBColor(50, 50, 50)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="475569"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    lines = code_text.strip('\r\n').split('\n')
    for idx, line in enumerate(lines):
        if idx == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        
        run = p.add_run(line if line else ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(8)
    return table

def create_table(doc, headers, rows, col_widths=None, caption=""):
    MAX_TABLE_WIDTH_INCHES = 6.25  # Standard A4 portrait printable width (8.27 - 1.18 - 0.79 = 6.30 in)

    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.bold = True
        run_cap.font.color.rgb = RGBColor(50, 50, 50)

    num_cols = len(headers)
    # Calculate and normalize column widths so sum never exceeds MAX_TABLE_WIDTH_INCHES
    if col_widths is None or len(col_widths) != num_cols:
        col_widths = [Inches(MAX_TABLE_WIDTH_INCHES / num_cols)] * num_cols
    else:
        total_w_in = sum(w.inches for w in col_widths)
        if total_w_in > MAX_TABLE_WIDTH_INCHES:
            scale = MAX_TABLE_WIDTH_INCHES / total_w_in
            col_widths = [Inches(w.inches * scale) for w in col_widths]

    # Typography and padding based on number of columns
    if num_cols >= 6:
        hdr_font_sz = Pt(9.5)
        body_font_sz = Pt(9.0)
        pad_top, pad_bot = 60, 60
        pad_lr = 70
        line_spacing = 1.05
    elif num_cols == 5:
        hdr_font_sz = Pt(10.0)
        body_font_sz = Pt(9.5)
        pad_top, pad_bot = 70, 70
        pad_lr = 85
        line_spacing = 1.10
    else:
        hdr_font_sz = Pt(10.5)
        body_font_sz = Pt(10.0)
        pad_top, pad_bot = 80, 80
        pad_lr = 110
        line_spacing = 1.15

    table = doc.add_table(rows=len(rows) + 1, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    # Set exact table width on tblPr
    tblPr = table._tbl.tblPr
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(MAX_TABLE_WIDTH_INCHES * 1440)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)

    # Repeat header row on every page (tblHeader)
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))

    # Style header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1A1C2C")
        set_cell_margins(hdr_cells[i], top=pad_top + 30, bottom=pad_bot + 30, left=pad_lr, right=pad_lr)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = hdr_font_sz
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Style body rows
    for r_idx, row in enumerate(rows):
        tr = table.rows[r_idx + 1]._tr.get_or_add_trPr()
        tr.append(OxmlElement('w:cantSplit'))  # Prevent row splitting across pages

        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            row_cells[c_idx].width = col_widths[c_idx]
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=pad_top, bottom=pad_bot, left=pad_lr, right=pad_lr)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = line_spacing
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if c_idx == 0 or len(str(val)) < 15:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = body_font_sz

    # Set column widths
    for i, col in enumerate(table.columns):
        col.width = col_widths[i]

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(8)
    return table

print("Word helper library initialized successfully!")
