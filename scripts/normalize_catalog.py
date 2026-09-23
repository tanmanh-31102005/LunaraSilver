"""Export the Phase 3 catalog from the original workbook and local media.

Run from the project root: python scripts/normalize_catalog.py
Requires openpyxl. The output is committed data; Laravel seeders do not read Excel.
"""

import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "docs/ProductData.xlsx"
OUTPUT = ROOT / "database/data/catalog.json"

CATEGORIES = [
    {"name": "Dây chuyền", "slug": "day-chuyen"},
    {"name": "Nhẫn", "slug": "nhan"},
    {"name": "Vòng tay", "slug": "vong-tay"},
    {"name": "Bộ trang sức", "slug": "bo-trang-suc"},
    {"name": "Set quà tặng", "slug": "set-qua-tang"},
]


def value(cell):
    result = cell.value
    return result.strip() if isinstance(result, str) else result


def money(cell, sku, field):
    raw = value(cell)
    if raw is None or raw == "":
        return None
    try:
        amount = Decimal(str(raw))
    except InvalidOperation as exc:
        raise ValueError(f"{sku} {field} {cell.parent.title}!{cell.coordinate}: invalid price {raw!r}") from exc
    if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
        raise ValueError(f"{sku} {field} {cell.parent.title}!{cell.coordinate}: invalid price {raw!r}")
    return format(amount, ".2f")


def media_path(folder, filename, sku, role, missing, required=False):
    path = ROOT / "media" / folder / filename
    if path.is_file():
        return path.relative_to(ROOT).as_posix()
    message = {"sku": sku, "role": role, "expected": path.relative_to(ROOT).as_posix()}
    if required:
        raise ValueError(f"{sku} image {role}: missing {message['expected']}")
    missing.append(message)
    return None


def collection_image(number, variant, sku, missing):
    candidates = sorted((ROOT / "media/Collection").glob(f"set {number}({variant}).*"))
    if len(candidates) > 1:
        raise ValueError(f"{sku} image ({variant}): ambiguous files {candidates}")
    if candidates:
        return candidates[0].relative_to(ROOT).as_posix()
    missing.append({"sku": sku, "role": "primary" if variant == 1 else "hover", "expected": f"media/Collection/set {number}({variant}).*"})
    return None


def components(raw, sku, source):
    # Only SKU references become relations. Non-retail additions remain in the
    # original composition text saved as short_description.
    pairs = re.findall(r"(?:([0-9]+)\s*[×x]\s*)?(LNS-(?:DC|NH|VT)[0-9]{3})", raw)
    if not pairs:
        raise ValueError(f"{sku} components {source}: no SKU found")
    result = []
    seen = set()
    for quantity, component in pairs:
        if component in seen or component == sku:
            raise ValueError(f"{sku} components {source}: duplicate or self-reference {component}")
        seen.add(component)
        result.append({"sku": component, "quantity": int(quantity) if quantity else 1})
    return result


def main():
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    singles = workbook["6.2 San pham"]
    sets = workbook["6.3 Set va qua tang"]
    products = []
    images = []
    bundles = []
    missing = []

    for row in range(4, 34):
        sku = value(singles.cell(row, 1))
        source = f"{singles.title}!{row}"
        category = value(singles.cell(row, 3))
        if category not in {item["name"] for item in CATEGORIES[:3]}:
            raise ValueError(f"{sku} category {source}: {category!r}")
        if value(singles.cell(row, 4)) != "single":
            raise ValueError(f"{sku} product_type {source}: expected single")
        stock = value(singles.cell(row, 7))
        if isinstance(stock, bool) or not isinstance(stock, int) or stock < 0:
            raise ValueError(f"{sku} stock_quantity {source}: invalid {stock!r}")
        regular = money(singles.cell(row, 5), sku, "regular_price")
        sale = money(singles.cell(row, 6), sku, "sale_price")
        if regular is None or (sale is not None and Decimal(sale) > Decimal(regular)):
            raise ValueError(f"{sku} price {source}: invalid regular/sale price")
        products.append({
            "source": source, "sku": sku, "name": value(singles.cell(row, 2)),
            "category_slug": next(item["slug"] for item in CATEGORIES if item["name"] == category),
            "product_type": "single", "regular_price": regular, "sale_price": sale,
            "stock_quantity": stock, "material": value(singles.cell(row, 8)),
            "stone": value(singles.cell(row, 9)), "weight": value(singles.cell(row, 10)),
            "size_info": value(singles.cell(row, 11)),
            "short_description": value(singles.cell(row, 12)),
            "description": value(singles.cell(row, 13)),
        })
        match = re.fullmatch(r"LNS-(DC|NH|VT)([0-9]{3})", sku)
        if not match:
            raise ValueError(f"{sku} sku {source}: unexpected single SKU")
        number = int(match.group(2))
        filename = f"{match.group(1).lower()}00{number if number == 10 else str(number)}.jpg"
        path = media_path("Product", filename, sku, "primary", missing, required=True)
        images.append({"sku": sku, "role": "primary", "path": path, "sort_order": 0})

    for row in list(range(4, 14)) + list(range(19, 24)):
        sku = value(sets.cell(row, 1))
        source = f"{sets.title}!{row}"
        is_gift = row >= 19
        number = row - 18 if is_gift else row - 3
        expected = f"LNS-{'GIFT' if is_gift else 'SET'}{number:03d}"
        if sku != expected:
            raise ValueError(f"{sku} sku {source}: expected {expected}")
        regular = money(sets.cell(row, 4), sku, "regular_price")
        if regular is None:
            raise ValueError(f"{sku} regular_price {source}: missing")
        composition = value(sets.cell(row, 3))
        products.append({
            "source": source, "sku": sku, "name": value(sets.cell(row, 2)),
            "category_slug": "set-qua-tang" if is_gift else "bo-trang-suc",
            "product_type": "gift" if is_gift else "collection",
            "regular_price": regular, "sale_price": None,
            # The workbook has no stock count for sets. Schema requires zero,
            # but bundle availability is derived from component stock instead.
            "stock_quantity": 0, "material": None, "stone": None, "weight": None,
            "size_info": None, "short_description": composition,
            "description": value(sets.cell(row, 10)),
        })
        bundles.append({"source": source, "sku": sku, "components": components(composition, sku, source)})
        if is_gift:
            path = media_path("Gift", f"set qua {number}.jpg.png", sku, "primary", missing, required=True)
            images.append({"sku": sku, "role": "primary", "path": path, "sort_order": 0})
        else:
            for variant, role in [(1, "primary"), (2, "hover")]:
                path = collection_image(number, variant, sku, missing)
                if path:
                    images.append({"sku": sku, "role": role, "path": path, "sort_order": variant - 1})

    skus = [product["sku"] for product in products]
    if len(skus) != len(set(skus)):
        raise ValueError("Duplicate product SKU in workbook")
    for bundle in bundles:
        for component in bundle["components"]:
            if component["sku"] not in skus:
                raise ValueError(f"{bundle['sku']} components {bundle['source']}: missing {component['sku']}")

    payload = {
        "source": "docs/ProductData.xlsx (sheets 6.2 San pham, 6.3 Set va qua tang) and media/Product, media/Collection, media/Gift",
        "categories": CATEGORIES, "products": products, "images": images,
        "bundles": bundles, "missing_images": missing,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Exported {len(products)} products, {len(images)} images, {len(bundles)} bundles; missing images: {missing}")


if __name__ == "__main__":
    main()
